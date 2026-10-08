#!/usr/bin/env python3
"""Genera el informe academico formal (Word) exigido en la FASE IX
(seccion 19) del enunciado:

"DISENO E IMPLEMENTACION DE INFRAESTRUCTURA COMO CODIGO PARA LA GESTION
MULTI-CLOUD DE 50 SUCURSALES MEDIANTE TERRAFORM Y PRUEBAS LOCALES"

Reutiliza:
  - El formato de portada (logo UTP, tabla Docente/Estudiantes, A4,
    margenes 2.5 cm, fuente Arial) extraido del documento de referencia
    Makip_Te_Crea_Proyecto_Actualizado.docx -- SOLO el formato, no su
    contenido (otro caso de negocio, fuera del alcance de este proyecto).
  - El contenido tecnico ya redactado en docs/architecture/*.md y
    docs/academic/*.md (via scripts/md_to_docx.py).
  - Los 8 diagramas generados en docs/diagrams/ (scripts/generate_diagrams.py)
    y las capturas reales de ejecucion en docs/evidence/screenshots/
    (scripts/render_terminal_screenshot.py).

Salida: docs/academic/Informe-IaC-50-Sucursales.docx

Uso:
    python scripts/generate_report.py
"""
from __future__ import annotations

import datetime
import pathlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor, Inches

from md_to_docx import add_markdown, add_figure

ROOT = pathlib.Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
DIAGRAMS = DOCS / "diagrams"
SCREENSHOTS = DOCS / "evidence" / "screenshots"
LOGO = ROOT / "docs" / "academic" / "assets" / "logo-utp.png"

MAROON = RGBColor(0x7A, 0x1D, 0x2A)

DOCENTE = "Eleodoro Orbegoso Carrera"
INTEGRANTES = [
    "Icochea Rodriguez, Thiago Paolo",
    "Gonzales Aguilar, Carlos Enrique Giussepe",
    "Huamani Pereira, Eddyson Cesar",
    "Torres Centeno, Emmanuel Misael",
    "Remuzgo Tovar, Huber Eduardo",
    "Quispe Saavedra, Karen Meylin",
]
TITULO = (
    "DISEÑO E IMPLEMENTACIÓN DE INFRAESTRUCTURA COMO CÓDIGO PARA LA GESTIÓN "
    "MULTI-CLOUD DE 50 SUCURSALES MEDIANTE TERRAFORM Y PRUEBAS LOCALES"
)


def set_base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), "Arial")

    for i in range(1, 5):
        h = doc.styles[f"Heading {i}"]
        h.font.name = "Arial"
        h.font.color.rgb = MAROON
        h.font.bold = True
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    doc.styles["Heading 1"].font.size = Pt(16)
    doc.styles["Heading 2"].font.size = Pt(13)
    doc.styles["Heading 3"].font.size = Pt(12)
    doc.styles["Heading 4"].font.size = Pt(11)


def set_page_setup(section) -> None:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def force_field_update_on_open(doc: Document) -> None:
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def add_page_number_footer(section) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE")


def build_cover(doc: Document) -> None:
    section = doc.sections[0]
    set_page_setup(section)
    section.different_first_page_header_footer = True

    if LOGO.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(7))

    def centered(text, size=12, bold=False, color=None, space_after=6):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        p.paragraph_format.space_after = Pt(space_after)
        return p

    centered("UNIVERSIDAD TECNOLÓGICA DEL PERÚ", size=15, bold=True, color=MAROON, space_after=2)
    centered("Facultad de Ingeniería — Curso: Servicios Cloud", size=12, space_after=24)

    centered(TITULO, size=16, bold=True, color=MAROON, space_after=30)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = 1
    table.style = "Table Grid"
    table.cell(0, 0).text = "Docente"
    table.cell(0, 1).text = DOCENTE
    table.cell(1, 0).text = "Integrantes"
    table.cell(1, 1).text = "\n".join(INTEGRANTES)
    for row in table.rows:
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs or [p.add_run("")]:
                    r.font.size = Pt(11)

    doc.add_paragraph()
    hoy = datetime.date.today()
    meses = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
             "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    centered(f"Lima, {hoy.day} de {meses[hoy.month - 1]} de {hoy.year}", size=11, space_after=2)
    centered(f"Año académico {hoy.year}", size=11)

    doc.add_section(WD_SECTION.NEW_PAGE)


def build_toc(doc: Document) -> None:
    doc.add_heading("ÍNDICE", level=1)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
    note = doc.add_paragraph()
    note.add_run(
        "(El índice se genera automáticamente al abrir el documento en Microsoft Word: "
        "si no se actualiza solo, usar clic derecho → Actualizar campo, o F9.)"
    ).italic = True
    doc.add_page_break()


def read(path: str) -> str:
    return (DOCS / path).read_text(encoding="utf-8")


def main() -> None:
    doc = Document()
    set_base_styles(doc)
    build_cover(doc)

    section2 = doc.sections[1]
    set_page_setup(section2)
    add_page_number_footer(section2)

    build_toc(doc)

    # RESUMEN --------------------------------------------------------------
    doc.add_heading("RESUMEN", level=1)
    add_markdown(doc, """
Este informe documenta el diseño, la implementación y la validación local de
una solución de Infraestructura como Código (IaC) con Terraform para
administrar **50 sucursales** de una empresa mediante un patrón común, sobre
**AWS y Google Cloud Platform**. Debido a que el proyecto **no cuenta con
credenciales de AWS ni de GCP**, la solución se valida íntegramente mediante
`terraform validate`, `terraform test` con `mock_provider`, un script de
validación de esquema en Python y dos laboratorios con recursos 100%
locales (`hashicorp/local`): uno para el ciclo de vida del Terraform State y
otro para la detección y reconciliación de Infrastructure Drift. Se
implementaron 8 módulos Terraform reutilizables (red, cómputo, seguridad y
monitoreo para cada proveedor), orquestados por un módulo compuesto
(`branch-stack`) que procesa las 50 sucursales a partir de un único archivo
de datos (`branches.yaml`), sin duplicar código. Se diseñó un pipeline
CI/CD común a ambos proveedores (GitHub Actions) con controles de formato,
validación, pruebas, análisis estático de seguridad (TFLint, Checkov) y
escaneo de secretos (gitleaks), y un procedimiento de aprobación mediante
Pull Request. Las 27 pruebas automatizadas de Terraform (`mock_provider`)
pasaron exitosamente, así como la validación de escalabilidad de las 50
sucursales y los dos laboratorios locales (State y Drift). El laboratorio
Docker de 2 sucursales (4 contenedores, HTTP, conectividad interna y
aislamiento de red) y el monitoreo local se ejecutaron realmente en GitHub
Actions —no en el equipo de desarrollo local, que no tiene Docker Desktop
operativo— con resultado exitoso y evidencia real capturada. El informe
distingue en todo momento lo que se ejecutó realmente de lo que queda
documentado como diseño para un despliegue empresarial futuro con
credenciales reales.

**Palabras clave**: Infraestructura como Código, Terraform, AWS, Google
Cloud Platform, CI/CD, Infrastructure Drift, Terraform State, DevOps.
""")
    doc.add_page_break()

    # 1. INTRODUCCION --------------------------------------------------------
    doc.add_heading("1. INTRODUCCIÓN", level=1)
    add_markdown(doc, """
La Infraestructura como Código (IaC) permite definir, versionar y automatizar
el aprovisionamiento de recursos cloud mediante archivos de configuración
declarativos, en lugar de procedimientos manuales. Herramientas como
Terraform de HashiCorp permiten administrar múltiples proveedores cloud
(AWS, GCP, Azure, entre otros) con una sintaxis común (HCL), manteniendo un
estado que refleja la infraestructura real administrada.

En escenarios empresariales con decenas de sucursales, sedes o unidades de
negocio que requieren una infraestructura tecnológica similar, la IaC
resuelve un problema central: aplicar el mismo patrón de red, cómputo,
seguridad y monitoreo de forma consistente, auditable y escalable, sin
copiar manualmente la configuración sucursal por sucursal. Este trabajo
aborda ese problema para el caso académico "cincuenta sucursales con un
patrón común", bajo la restricción adicional de no contar con credenciales
reales de AWS ni GCP, lo que exige una estrategia de validación basada en
simulación, pruebas locales y proveedores simulados (`mock_provider`).
""")

    # 2. PLANTEAMIENTO DEL PROBLEMA ------------------------------------------
    doc.add_heading("2. PLANTEAMIENTO DEL PROBLEMA", level=1)
    add_markdown(doc, read("academic/analisis-requisitos.md"), skip_first_h1=True)

    # 3. OBJETIVOS (reuse part of analisis-requisitos is above; add explicit objetivos again as its own numbered section per outline)
    doc.add_heading("3. OBJETIVOS", level=1)
    add_markdown(doc, """
## Objetivo general

Diseñar, implementar y validar localmente una solución de Infraestructura
como Código con Terraform que aprovisione el patrón común de 50 sucursales
en AWS y GCP, de forma reutilizable, segura, probada y documentada, sin
requerir credenciales cloud para su demostración.

## Objetivos específicos

1. Diseñar una arquitectura multi-cloud con interfaz de configuración común
   entre AWS y GCP.
2. Implementar módulos Terraform reutilizables para ambos proveedores.
3. Demostrar que la solución escala a 50 sucursales sin duplicar código.
4. Implementar una estrategia de pruebas sin credenciales cloud.
5. Diseñar la gestión del Terraform State (estrategia empresarial +
   laboratorio local funcional).
6. Implementar un pipeline CI/CD común con controles de seguridad.
7. Demostrar el ciclo completo de detección y reconciliación de
   Infrastructure Drift.
8. Documentar el proyecto en formato académico IEEE y publicarlo en GitHub.
""")

    # 4. MARCO TEORICO --------------------------------------------------------
    doc.add_heading("4. MARCO TEÓRICO", level=1)
    add_markdown(doc, """
## 4.1 Infraestructura como Código

La Infraestructura como Código es la práctica de gestionar y aprovisionar
infraestructura computacional mediante archivos de definición legibles por
máquina, en lugar de configuración manual o herramientas interactivas [1].

## 4.2 Terraform

Terraform es una herramienta de IaC de código abierto desarrollada por
HashiCorp que usa el lenguaje declarativo HCL (HashiCorp Configuration
Language) para definir recursos de múltiples proveedores cloud, manteniendo
un archivo de estado que mapea la configuración declarada con los recursos
reales [2]. Desde la versión 1.6, incorpora `terraform test`, un framework
de pruebas nativo que soporta `mock_provider` para simular proveedores sin
realizar llamadas reales a sus APIs [3].

## 4.3 AWS

Amazon Web Services (AWS) es la plataforma cloud de Amazon, que ofrece
servicios de cómputo (EC2), redes (VPC), seguridad (Security Groups, IAM) y
observabilidad (CloudWatch), entre otros [4].

## 4.4 Google Cloud Platform

Google Cloud Platform (GCP) es la plataforma cloud de Google, con servicios
equivalentes: Compute Engine, VPC Network, Firewall Rules y Cloud
Monitoring [5].

## 4.5 CI/CD

La integración y entrega continuas (CI/CD) automatizan la construcción,
prueba y despliegue de software mediante pipelines declarativos. GitHub
Actions es la plataforma de CI/CD nativa de GitHub, basada en workflows
YAML disparados por eventos del repositorio [6].

## 4.6 Terraform State

El estado de Terraform es un archivo (por defecto JSON) que registra los
recursos que Terraform administra y sus atributos, permitiendo calcular
diferencias entre la configuración declarada y la infraestructura real [7].
Para equipos, HashiCorp recomienda backends remotos con bloqueo
(locking) para evitar condiciones de carrera entre ejecuciones concurrentes [8].

## 4.7 Infrastructure Drift

El "drift" de infraestructura ocurre cuando el estado real de un recurso
diverge de lo declarado en el código, típicamente por cambios manuales
realizados fuera de la herramienta de IaC. Terraform detecta el drift
mediante `terraform plan`, que compara el estado almacenado (opcionalmente
refrescado contra el proveedor real) con la configuración declarada [9].

## 4.8 Seguridad Cloud

Las prácticas de seguridad cloud relevantes para este proyecto incluyen el
principio de mínimo privilegio en IAM, el cifrado de datos en reposo y en
tránsito, la gestión de secretos fuera del control de versiones, y el
análisis estático de configuraciones IaC (p. ej., con Checkov) para detectar
configuraciones inseguras antes del despliegue [10].
""")

    # 5. ANALISIS DE REQUISITOS ------------------------------------------
    doc.add_heading("5. ANÁLISIS DE REQUISITOS", level=1)
    add_markdown(doc, read("academic/matriz-trazabilidad.md"), skip_first_h1=True)

    # 6. ARQUITECTURA PROPUESTA ------------------------------------------
    doc.add_heading("6. ARQUITECTURA PROPUESTA", level=1)
    add_markdown(doc, read("architecture/overview.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-01-arquitectura-general.png", "Figura 1. Arquitectura general de las 50 sucursales.")
    add_markdown(doc, read("architecture/aws-architecture.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-02-arquitectura-aws.png", "Figura 2. Arquitectura AWS por sucursal.")
    add_markdown(doc, read("architecture/gcp-architecture.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-03-arquitectura-gcp.png", "Figura 3. Arquitectura GCP por sucursal.")

    # 7. IMPLEMENTACION ----------------------------------------------------
    doc.add_heading("7. IMPLEMENTACIÓN", level=1)
    add_markdown(doc, """
## 7.1 Estructura del repositorio

El repositorio `terraform-iac-50-sucursales` organiza el código en
`terraform/` (módulos y entornos), `lab/` (laboratorios locales), `config/`
(datos de las sucursales), `scripts/` (automatización y generación de
reportes), `.github/workflows/` (pipeline CI/CD) y `docs/` (arquitectura,
documentación académica y evidencia).
""")
    add_figure(doc, DIAGRAMS / "fig-04-estructura-modulos.png", "Figura 4. Estructura de módulos Terraform.")
    add_markdown(doc, """
## 7.2 Módulos Terraform

Se implementaron 8 módulos base (`network`, `compute`, `security`,
`monitoring` para AWS y para GCP) más un módulo orquestador
(`branch-stack`), todos con `main.tf`, `variables.tf`, `outputs.tf`,
`versions.tf`, validaciones de entrada (`validation` blocks) y un `README.md`
con ejemplos de uso.

## 7.3 Configuración de sucursales

`config/branches/branches.yaml` es la fuente de verdad de las 50 sucursales,
generada de forma determinista por `scripts/generate_branches.py`. Cada
sucursal define: identificador, nombre, proveedor, región, entorno,
configuración de red, dos servidores, reglas de seguridad, monitoreo y
etiquetas (ver `config/branches/README.md`).

## 7.4 Redes, servidores, seguridad y monitoreo

Ver secciones 6.2 y 6.3 (arquitectura AWS/GCP) para el detalle de los
recursos creados por sucursal.
""")

    # 8. GESTION DEL ESTADO -------------------------------------------------
    doc.add_heading("8. GESTIÓN DEL ESTADO", level=1)
    add_markdown(doc, read("architecture/terraform-state-strategy.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-06-terraform-state.png", "Figura 6. Gestión del Terraform State.")

    # 9. PIPELINE CI/CD -------------------------------------------------
    doc.add_heading("9. PIPELINE CI/CD", level=1)
    add_markdown(doc, read("architecture/cicd-pipeline.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-05-flujo-cicd.png", "Figura 5. Flujo CI/CD común a AWS y GCP.")

    # 10. SEGURIDAD -------------------------------------------------------
    doc.add_heading("10. SEGURIDAD", level=1)
    add_markdown(doc, read("architecture/secrets-management.md"), skip_first_h1=True)

    # 11. INFRASTRUCTURE DRIFT ----------------------------------------------
    doc.add_heading("11. INFRASTRUCTURE DRIFT", level=1)
    add_markdown(doc, read("../lab/drift/README.md"), skip_first_h1=True)
    add_figure(doc, DIAGRAMS / "fig-07-procedimiento-drift.png", "Figura 7. Procedimiento de Infrastructure Drift.")
    add_figure(doc, SCREENSHOTS / "fig-drift-deteccion.png", "Figura 9a. Evidencia real: detección de drift (terraform plan).")
    add_figure(doc, SCREENSHOTS / "fig-drift-reconciliacion.png", "Figura 9b. Evidencia real: verificación tras la reconciliación.")

    # 12. PRUEBAS LOCALES -------------------------------------------------
    doc.add_heading("12. PRUEBAS LOCALES", level=1)
    add_markdown(doc, read("academic/plan-de-pruebas.md"), skip_first_h1=True)
    add_figure(doc, SCREENSHOTS / "fig-terraform-validate.png", "Figura 9c. Evidencia real: terraform validate en los 3 entornos.")
    add_figure(doc, SCREENSHOTS / "fig-terraform-test.png", "Figura 9d. Evidencia real: terraform test (27 passed, 0 failed).")
    add_figure(doc, SCREENSHOTS / "fig-scalability-validation.png", "Figura 9e. Evidencia real: validación de las 50 sucursales.")
    add_figure(doc, DIAGRAMS / "fig-08-laboratorio-docker.png", "Figura 8. Arquitectura del laboratorio local (Docker).")
    add_figure(doc, SCREENSHOTS / "fig-docker-lab.png", "Figura 9f. Evidencia real: laboratorio Docker ejecutado en GitHub Actions (4 contenedores healthy, HTTP 200, conectividad interna y aislamiento de red confirmados).")

    # 13. RESULTADOS -------------------------------------------------
    doc.add_heading("13. RESULTADOS", level=1)
    add_markdown(doc, """
De 16 casos de prueba planificados (CP01-CP16), los 16 se ejecutaron y
pasaron con evidencia real. Las 27 pruebas de `terraform test` con
`mock_provider` pasaron sin fallos tras corregir un error real detectado
durante la primera ejecución (indexación inválida sobre un atributo de tipo
`set` en el módulo `aws/network`; ver CP02 en el plan de pruebas). La
validación de escalabilidad confirmó 50 sucursales (25 AWS / 25 GCP), 40 en
producción, 6 en staging y 4 en desarrollo, sin identificadores duplicados.
El laboratorio Docker de 2 sucursales (4 contenedores, HTTP, conectividad
interna, aislamiento de red) y el monitoreo local se ejecutaron realmente
en un runner `ubuntu-latest` de GitHub Actions —que trae Docker Engine y
Docker Compose preinstalados— ya que el equipo de desarrollo local no tiene
Docker Desktop operativo (instalado, pero sin poder iniciar sin reiniciar
Windows, reinicio que se decidió no realizar). El pipeline de CI/CD
detectó, en su primera ejecución real, dos problemas genuinos (formato sin
aplicar y advertencias de TFLint en los módulos GCP; un paso de
inicialización de backend incompleto en el workflow de Drift), ambos
corregidos y verificados con una segunda ejecución exitosa.

## Limitaciones

- No se validó el comportamiento contra la API real de AWS ni GCP (sin
  credenciales).
- El laboratorio Docker y el monitoreo local se validaron en un runner de
  GitHub Actions, no en el equipo de desarrollo Windows (que tiene Docker
  Desktop instalado pero no operativo, por la razón descrita arriba). El
  código ejecutado es idéntico en ambos casos; solo cambia el entorno.
""")

    # 14. DISCUSION -------------------------------------------------
    doc.add_heading("14. DISCUSIÓN", level=1)
    add_markdown(doc, """
El diseño empresarial (backends remotos S3/GCS, IAM de mínimo privilegio,
GitHub OIDC, despliegue real tras aprobación) y el laboratorio local
ejecutado en este proyecto comparten el mismo código Terraform de los
módulos: la diferencia está exclusivamente en el backend de estado, la
autenticación y si el pipeline ejecuta `apply` contra un proveedor real o
simulado. Esto significa que migrar de la demostración académica a un
despliegue empresarial real no requiere reescribir los módulos: requiere
configurar backends remotos, credenciales sin claves estáticas (OIDC) y
habilitar los jobs de `apply` actualmente deshabilitados o documentados
como futuros (ver `docs/architecture/cicd-pipeline.md`, sección "Habilitar
el despliegue real en el futuro").
""")

    # 15. CONCLUSIONES -------------------------------------------------
    doc.add_heading("15. CONCLUSIONES", level=1)
    add_markdown(doc, """
1. Es posible diseñar, implementar y validar una solución de IaC
   multi-cloud para 50 sucursales sin credenciales reales, usando
   `mock_provider`, validación de esquema y laboratorios con recursos
   locales, cumpliendo el objetivo general del proyecto.
2. La interfaz común entre los módulos AWS y GCP (campos lógicos `size`,
   `rules`, `servers`) permite agregar sucursales como un cambio de datos
   (`branches.yaml`), sin duplicar ni modificar código Terraform,
   cumpliendo el requisito de extensibilidad (RNF-01).
3. `terraform test` con `mock_provider` detectó un error real de la
   configuración (indexación sobre un `set`) antes de cualquier intento de
   despliegue, validando el valor de las pruebas automatizadas incluso sin
   credenciales cloud.
4. El mecanismo de detección y reconciliación de Infrastructure Drift
   (`terraform plan`, `plan -refresh-only`, `apply`) es el mismo
   independientemente del proveedor; se demostró end-to-end con un recurso
   local real.
5. La principal limitación del proyecto es la imposibilidad de validar el
   comportamiento contra las APIs reales de AWS/GCP; el laboratorio Docker,
   que inicialmente tampoco podía ejecutarse localmente, se completó con
   evidencia real usando un runner de GitHub Actions en lugar de simular su
   resultado.
""")

    # 16. RECOMENDACIONES -------------------------------------------------
    doc.add_heading("16. RECOMENDACIONES", level=1)
    add_markdown(doc, """
1. Para una demostración presencial, ejecutar también el laboratorio Docker en
   un equipo con Docker Desktop disponible, para validar el patrón de
   aislamiento de red con evidencia adicional.
2. Configurar backends remotos (S3/GCS) y autenticación OIDC antes de
   habilitar cualquier job de `terraform apply` en el pipeline CI/CD.
3. Incorporar `terraform plan` contra el backend remoto real como paso de
   revisión humana obligatorio antes de cualquier `apply` en producción.
4. Evaluar alias de provider por región en AWS si el negocio real requiere
   sucursales AWS en más de una región simultáneamente.
5. Extender la suite de `terraform test` con casos adicionales a medida que
   se agreguen nuevos tipos de recursos (ej. balanceadores de carga,
   bases de datos gestionadas) a los módulos.
""")

    # 17. REFERENCIAS BIBLIOGRAFICAS -------------------------------------------------
    doc.add_heading("17. REFERENCIAS BIBLIOGRÁFICAS", level=1)
    refs = [
        'HashiCorp, "What is Infrastructure as Code with Terraform?," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/intro. [Accessed: Oct. 2026].',
        'HashiCorp, "Terraform Language Documentation," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/language. [Accessed: Oct. 2026].',
        'HashiCorp, "Tests - Terraform CLI Documentation," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/language/tests. [Accessed: Oct. 2026].',
        'Amazon Web Services, "AWS Documentation," AWS. [Online]. Available: https://docs.aws.amazon.com/. [Accessed: Oct. 2026].',
        'Google LLC, "Google Cloud Documentation," Google Cloud. [Online]. Available: https://cloud.google.com/docs. [Accessed: Oct. 2026].',
        'GitHub Inc., "GitHub Actions Documentation," GitHub Docs. [Online]. Available: https://docs.github.com/actions. [Accessed: Oct. 2026].',
        'HashiCorp, "State - Terraform CLI Documentation," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/language/state. [Accessed: Oct. 2026].',
        'HashiCorp, "Backend Configuration," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/language/backend. [Accessed: Oct. 2026].',
        'HashiCorp, "Resource Drift," HashiCorp Developer. [Online]. Available: https://developer.hashicorp.com/terraform/language/state/resource-drift. [Accessed: Oct. 2026].',
        'Bridgecrew (Prisma Cloud), "Checkov Documentation," checkov.io. [Online]. Available: https://www.checkov.io/. [Accessed: Oct. 2026].',
        'Docker Inc., "Docker Compose Documentation," Docker Docs. [Online]. Available: https://docs.docker.com/compose/. [Accessed: Oct. 2026].',
        'Gitleaks, "Gitleaks Documentation," GitHub. [Online]. Available: https://github.com/gitleaks/gitleaks. [Accessed: Oct. 2026].',
        'Terraform Linters, "TFLint Documentation," GitHub. [Online]. Available: https://github.com/terraform-linters/tflint. [Accessed: Oct. 2026].',
    ]
    for i, ref in enumerate(refs, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        p.add_run(f"[{i}] {ref}")

    # 18. ANEXOS -------------------------------------------------
    doc.add_heading("18. ANEXOS", level=1)
    add_markdown(doc, """
## Anexo A — Comandos principales

```
./scripts/validate.ps1
./scripts/test.ps1
cd lab/local-terraform; ./run_state_demo.ps1
cd lab/drift; ./run_drift_demo.ps1
cd lab/docker; ./start.ps1; ./test.ps1; ./stop.ps1
```

## Anexo B — Enlace al repositorio GitHub

https://github.com/ThiagoIcochea/ServiciosCloudIAC

## Anexo C — Guía de ejecución local completa

Ver `docs/user-guide/README.md` en el repositorio.
""")

    out_dir = DOCS / "academic"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "Informe-IaC-50-Sucursales.docx"
    force_field_update_on_open(doc)
    doc.save(out_path)
    print(f"Informe generado: {out_path}")


if __name__ == "__main__":
    main()
