"""Conversor ligero de un subconjunto propio de Markdown a python-docx.

No es un conversor Markdown general: solo entiende la sintaxis que se usa
consistentemente en los documentos de docs/academic y docs/architecture de
este proyecto (encabezados #/##/###, tablas con |, listas con - o 1., negrita
**texto**, codigo en linea `texto`, bloques de codigo con ```, citas con >).
Los bloques ```mermaid``` se omiten (los diagramas equivalentes ya existen
como imagenes PNG en docs/diagrams, insertadas aparte por el script que
ensambla el informe).
"""
from __future__ import annotations

import re

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
CODE_RE = re.compile(r"`([^`]+)`")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")

HEADING_COLOR = RGBColor(0x7A, 0x1D, 0x2A)


def _add_inline_runs(paragraph, text: str) -> None:
    """Agrega texto a un paragraph interpretando **negrita**, `codigo` y
    [texto](link) (este ultimo se aplana a 'texto (link)')."""
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)

    pos = 0
    tokens: list[tuple[str, str]] = []
    pattern = re.compile(r"(\*\*(?:.+?)\*\*|`[^`]+`)")
    for part in pattern.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            tokens.append(("bold", part[2:-2]))
        elif part.startswith("`") and part.endswith("`"):
            tokens.append(("code", part[1:-1]))
        else:
            tokens.append(("text", part))

    for kind, value in tokens:
        run = paragraph.add_run(value)
        if kind == "bold":
            run.bold = True
        elif kind == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)


def _parse_table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        if re.match(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)+\|?\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def _style_table(table) -> None:
    table.style = "Light Grid Accent 2"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True


def add_markdown(doc, md_text: str, heading_base: int = 1, skip_first_h1: bool = False) -> None:
    lines = md_text.replace("\r\n", "\n").split("\n")
    i = 0
    n = len(lines)
    first_h1_seen = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Bloques de codigo ```lang ... ```
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            body = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1  # cerrar ```
            if lang == "mermaid":
                continue  # el diagrama equivalente ya se inserta como figura
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            shade = p._p.get_or_add_pPr()
            shd = shade.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:fill"): "F2F2F2"})
            shade.append(shd)
            for j, code_line in enumerate(body):
                run = p.add_run(code_line if j == 0 else "\n" + code_line)
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            continue

        # Encabezados
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and skip_first_h1 and not first_h1_seen:
                first_h1_seen = True
                i += 1
                continue
            doc.add_heading(text, level=min(heading_base + level - 1, 9))
            i += 1
            continue

        # Tablas
        if stripped.startswith("|"):
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = _parse_table_rows(table_lines)
            if rows:
                ncols = len(rows[0])
                table = doc.add_table(rows=0, cols=ncols)
                for row_cells in rows:
                    row = table.add_row()
                    for ci in range(ncols):
                        text = row_cells[ci] if ci < len(row_cells) else ""
                        cell_p = row.cells[ci].paragraphs[0]
                        _add_inline_runs(cell_p, text)
                _style_table(table)
            continue

        # Listas
        m_bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        m_num = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m_bullet or m_num:
            style = "List Bullet" if m_bullet else "List Number"
            text = (m_bullet or m_num).group(1)
            p = doc.add_paragraph(style=style)
            _add_inline_runs(p, text)
            i += 1
            continue

        # Citas
        if stripped.startswith(">"):
            text = stripped.lstrip(">").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run = p.add_run(text)
            run.italic = True
            i += 1
            continue

        # Separador horizontal
        if re.match(r"^-{3,}$", stripped):
            i += 1
            continue

        # Parrafo normal
        p = doc.add_paragraph()
        _add_inline_runs(p, stripped)
        i += 1


def add_figure(doc, image_path, caption: str, width_inches: float = 6.0) -> None:
    from docx.shared import Inches

    doc.add_picture(str(image_path), width=Inches(width_inches))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
